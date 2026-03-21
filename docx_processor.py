"""Extract and replace text runs in .docx files while preserving formatting.

Handles text in:
- Body paragraphs
- Table cells (body and inside headers/footers)
- Headers and footers (including text inside shapes, text boxes, WordArt)
"""

from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table, _Cell

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


class XmlTextRun:
    """Lightweight wrapper around an XML w:r element for direct text replacement.

    Used for text inside shapes/textboxes where python-docx doesn't expose Run objects.
    """

    def __init__(self, r_element, t_element):
        self._r_element = r_element
        self._t_element = t_element

    @property
    def text(self):
        return self._t_element.text or ""

    @text.setter
    def text(self, value):
        self._t_element.text = value


def _extract_runs_from_paragraph(paragraph: Paragraph, location: str):
    """Yield (run, original_text, location_info) for each non-empty run in a paragraph."""
    for run in paragraph.runs:
        if run.text.strip():
            yield (run, run.text, location)


def _extract_runs_from_cell(cell: _Cell, table_idx: int, row_idx: int, col_idx: int):
    """Yield runs from all paragraphs inside a table cell."""
    for para in cell.paragraphs:
        loc = f"table[{table_idx}].row[{row_idx}].col[{col_idx}]"
        yield from _extract_runs_from_paragraph(para, loc)


def _extract_runs_from_table(table: Table, table_idx):
    """Yield runs from all cells in a table."""
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            yield from _extract_runs_from_cell(cell, table_idx, row_idx, col_idx)


def _extract_runs_from_header_footer_deep(hf, label: str):
    """Extract ALL text runs from header/footer XML, including inside shapes/textboxes.

    Goes to the XML level to find every w:r that contains w:t with non-empty text.
    This catches text in standard paragraphs, text boxes, shapes, WordArt, etc.
    """
    hf_element = hf._element

    # Standard paragraphs and tables (python-docx objects)
    for para in hf.paragraphs:
        yield from _extract_runs_from_paragraph(para, label)
    for table_idx, table in enumerate(hf.tables):
        yield from _extract_runs_from_table(table, f"{label}.table[{table_idx}]")

    # Deep XML scan for text inside shapes, textboxes, etc.
    # Find ALL w:r elements in the header/footer XML
    known_r_elements = set()
    for para in hf.paragraphs:
        for run in para.runs:
            known_r_elements.add(run._element)

    # Also collect known table runs
    for table in hf.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        known_r_elements.add(run._element)

    # Find all w:r elements with w:t children that have text
    for r_elem in hf_element.findall(f".//{{{W_NS}}}r"):
        if r_elem in known_r_elements:
            continue  # Already handled via python-docx objects

        t_elems = r_elem.findall(f"{{{W_NS}}}t")
        for t_elem in t_elems:
            if t_elem.text and t_elem.text.strip():
                wrapper = XmlTextRun(r_elem, t_elem)
                yield (wrapper, t_elem.text, f"{label}.shape")
                break  # One wrapper per w:r


def _extract_runs_from_header_footer(section, section_idx: int):
    """Yield runs from all headers and footers of a section."""
    for hf, label in [
        (section.header, f"section[{section_idx}].header"),
        (section.footer, f"section[{section_idx}].footer"),
        (section.first_page_header, f"section[{section_idx}].first_page_header"),
        (section.first_page_footer, f"section[{section_idx}].first_page_footer"),
        (section.even_page_header, f"section[{section_idx}].even_page_header"),
        (section.even_page_footer, f"section[{section_idx}].even_page_footer"),
    ]:
        if hf is not None:
            yield from _extract_runs_from_header_footer_deep(hf, label)


def extract_all_runs(doc: Document):
    """Extract all text runs from a document (body, tables, headers, footers, shapes).

    Returns a list of tuples: (run_object, original_text, location_description)
    """
    runs = []

    # Body paragraphs
    for para in doc.paragraphs:
        runs.extend(_extract_runs_from_paragraph(para, "body"))

    # Body tables
    for table_idx, table in enumerate(doc.tables):
        runs.extend(_extract_runs_from_table(table, table_idx))

    # Headers and footers from all sections (deep scan)
    for section_idx, section in enumerate(doc.sections):
        runs.extend(_extract_runs_from_header_footer(section, section_idx))

    return runs


def apply_translations(run_items, translations):
    """Replace text in runs with translated text.

    Args:
        run_items: list of (run_obj, original_text, location) tuples
        translations: list of translated strings, same order as run_items
    """
    for (run, original, _loc), translated in zip(run_items, translations):
        if translated is not None and translated.strip():
            # Preserve leading/trailing whitespace from original
            leading = len(original) - len(original.lstrip())
            trailing = len(original) - len(original.rstrip())
            prefix = original[:leading]
            suffix = original[-trailing:] if trailing > 0 else ""
            run.text = prefix + translated.strip() + suffix
